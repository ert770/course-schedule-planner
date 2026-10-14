# Write a review workbook from a JSON payload produced by a Node script.
# Payload: {"output": path, "sheets": [{"name", "columns", "rows"}], "widths": [...],
#           "decisionColumn": 1-based column index or null, "decisionOptions": [...]}
# Text stays in the UTF-8 JSON payload so this file is ASCII only.
import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


def main(payload_path):
    with open(payload_path, encoding="utf-8") as handle:
        payload = json.load(handle)
    book = Workbook()
    book.remove(book.active)
    header_fill = PatternFill("solid", fgColor="DDE7F3")
    decision_fill = PatternFill("solid", fgColor="FFF4CC")
    widths = payload.get("widths") or []
    decision_column = payload.get("decisionColumn")
    options = payload.get("decisionOptions") or []

    for spec in payload["sheets"]:
        sheet = book.create_sheet(spec["name"][:31])
        sheet.append(spec["columns"])
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
        for row in spec["rows"]:
            sheet.append(row)
        sheet.freeze_panes = "A2"
        is_info = len(spec["columns"]) == 2
        if is_info:
            sheet.column_dimensions["A"].width = 14
            sheet.column_dimensions["B"].width = 110
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
            continue
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        last_row = sheet.max_row
        if last_row > 1:
            sheet.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(spec["columns"])), last_row)
        if decision_column and options and last_row > 1:
            letter = get_column_letter(decision_column)
            validation = DataValidation(type="list", formula1='"%s"' % ",".join(options), allow_blank=True)
            sheet.add_data_validation(validation)
            validation.add("%s2:%s%d" % (letter, letter, last_row))
            for row in range(2, last_row + 1):
                sheet["%s%d" % (letter, row)].fill = decision_fill

    book.save(payload["output"])
    print("written", payload["output"], [s["name"] for s in payload["sheets"]], [len(s["rows"]) for s in payload["sheets"]])


if __name__ == "__main__":
    main(sys.argv[1])
