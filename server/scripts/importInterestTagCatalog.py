#!/usr/bin/env python3
"""Build the runtime rag_tag taxonomy catalog from the approved workbook.

Requires openpyxl in the local maintenance environment. The production server
loads the generated JSON and does not depend on Python or openpyxl.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import sys
import unicodedata
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
DEFAULT_WORKBOOK = ROOT / "docs/CHANGE_REPORTS/2026-10-08-rag-tag-exclusions-semantic-merge-approved.xlsx"
DEFAULT_ALIASES = ROOT / "server/src/data/interestTagAliases.json"
DEFAULT_OUTPUT = ROOT / "server/src/data/interestTagCatalog.json"
CATALOG_VERSION = "rag-tag-catalog-2026-10-08-v1"
CLASSIFICATION_SHEET_PREFIX = "分類標籤_"
EXPECTED_CLASSIFICATION_SHEETS = 31
EXPECTED_ASSIGNMENT_ROWS = 15113


def clean_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def comparison_key(value: object) -> str:
    text = unicodedata.normalize("NFKC", clean_text(value)).casefold()
    return "".join(
        character
        for character in text
        if not character.isspace()
        and unicodedata.category(character) != "Pd"
        and character != "\u2212"
    )


def stable_id(prefix: str, key: str) -> str:
    digest = hashlib.sha256(key.encode("utf-8")).hexdigest()[:20]
    return f"{prefix}_{digest}"


def category_headers(workbook):
    sheets = [sheet for sheet in workbook.worksheets if sheet.title.startswith(CLASSIFICATION_SHEET_PREFIX)]
    sheets.sort(key=lambda sheet: sheet.title)
    if len(sheets) != EXPECTED_CLASSIFICATION_SHEETS:
        raise ValueError(
            f"Expected {EXPECTED_CLASSIFICATION_SHEETS} classification pages, found {len(sheets)}"
        )

    required_headers = {
        "主分類": "mainCategory",
        "子分類（拆分後）": "subcategory",
        "原始標籤": "rawTag",
    }
    for sheet in sheets:
        header_row_index = None
        header_values = None
        for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            values = [clean_text(value) for value in row]
            if "主分類" in values and "原始標籤" in values:
                header_row_index = row_index
                header_values = values
                break
            if row_index >= 12:
                break
        if header_row_index is None or header_values is None:
            raise ValueError(f"Could not find classification headers in {sheet.title}")

        column_indexes = {name: header_values.index(name) for name in required_headers if name in header_values}
        missing = sorted(set(required_headers) - set(column_indexes))
        if missing:
            raise ValueError(f"Missing required columns in {sheet.title}: {', '.join(missing)}")
        yield sheet, header_row_index, required_headers, column_indexes


def load_aliases(path: Path):
    payload = json.loads(path.read_text(encoding="utf-8"))
    aliases = payload.get("aliases", {})
    if not isinstance(aliases, dict):
        raise ValueError("Alias file must contain an aliases object")
    alias_index = {}
    for alias, canonical in aliases.items():
        alias_key = comparison_key(alias)
        canonical_label = clean_text(canonical)
        if alias_key and canonical_label:
            alias_index[alias_key] = canonical_label
    return payload, alias_index


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--aliases", type=Path, default=DEFAULT_ALIASES)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()

    _, alias_index = load_aliases(args.aliases)
    workbook = load_workbook(args.workbook, read_only=True, data_only=True)

    main_categories = {}
    subcategories = {}
    canonical_tags = {}
    raw_tag_keys = set()
    raw_tag_paths = defaultdict(set)
    assignment_count = 0
    category_sheet_count = 0

    for sheet, header_row_index, columns, column_indexes in category_headers(workbook):
        category_sheet_count += 1
        for row_number, row_values in enumerate(
            sheet.iter_rows(min_row=header_row_index + 1, values_only=True),
            start=header_row_index + 1,
        ):
            row = {field: clean_text(row_values[index] if index < len(row_values) else None)
                   for header, field in columns.items()
                   for index in [column_indexes[header]]}
            if not any(row.values()):
                continue

            main_name = row["mainCategory"]
            subcategory_name = row["subcategory"]
            raw_tag = row["rawTag"]
            if not main_name or not subcategory_name or not raw_tag:
                raise ValueError(f"Incomplete taxonomy row at {sheet.title}:{row_number}")

            main_key = comparison_key(main_name)
            subcategory_key = comparison_key(subcategory_name)
            raw_key = comparison_key(raw_tag)
            canonical_label = alias_index.get(raw_key, raw_tag)
            canonical_key = comparison_key(canonical_label)
            if not main_key or not subcategory_key or not raw_key or not canonical_key:
                raise ValueError(f"Empty normalized taxonomy value at {sheet.title}:{row_number}")

            main_id = stable_id("main", main_key)
            path_key = f"{main_key}\u0000{subcategory_key}"
            subcategory_id = stable_id("sub", path_key)
            main_categories.setdefault(main_id, {"id": main_id, "name": main_name})
            subcategories.setdefault(subcategory_id, {
                "id": subcategory_id,
                "mainCategoryId": main_id,
                "name": subcategory_name,
            })

            tag = canonical_tags.setdefault(canonical_key, {
                "id": stable_id("tag", canonical_key),
                "name": canonical_label,
                "normalizedName": canonical_key,
                "categoryAssignments": {},
                "eligibility": {
                    "status": "pending_post_alias_course_recount",
                    "interestLearningEligible": None,
                    "crossCourseMatchEligible": None,
                    "exclusionReason": None,
                },
            })
            # An approved alias target is the canonical display name even if its
            # row appears after an alias row in the workbook.
            if canonical_label != raw_tag or tag["name"] == raw_tag:
                tag["name"] = canonical_label

            assignment_key = f"{main_id}\u0000{subcategory_id}"
            assignment = tag["categoryAssignments"].setdefault(assignment_key, {
                "mainCategoryId": main_id,
                "subcategoryId": subcategory_id,
                "sourceRows": [],
            })
            assignment["sourceRows"].append({
                "rawTag": raw_tag,
                "worksheet": sheet.title,
                "row": row_number,
            })

            raw_tag_keys.add(raw_tag)
            raw_tag_paths[raw_tag].add(path_key)
            assignment_count += 1

    if category_sheet_count != EXPECTED_CLASSIFICATION_SHEETS:
        raise ValueError(f"Only processed {category_sheet_count} classification pages")
    if assignment_count != EXPECTED_ASSIGNMENT_ROWS:
        raise ValueError(
            f"Expected {EXPECTED_ASSIGNMENT_ROWS} taxonomy assignments, found {assignment_count}"
        )

    alias_targets_missing = sorted({
        canonical
        for canonical in alias_index.values()
        if comparison_key(canonical) not in canonical_tags
    })

    serialized_tags = []
    for tag in canonical_tags.values():
        assignments = []
        for assignment in tag["categoryAssignments"].values():
            assignments.append({
                "mainCategoryId": assignment["mainCategoryId"],
                "subcategoryId": assignment["subcategoryId"],
                "sourceRows": sorted(
                    assignment["sourceRows"],
                    key=lambda item: (comparison_key(item["rawTag"]), item["rawTag"], item["worksheet"], item["row"]),
                ),
            })
        assignments.sort(key=lambda item: (
            item["mainCategoryId"], item["subcategoryId"]
        ))
        serialized_tags.append({
            **{key: value for key, value in tag.items() if key != "categoryAssignments"},
            "categoryAssignments": assignments,
        })
    serialized_tags.sort(key=lambda item: (item["normalizedName"], item["id"]))

    catalog = {
        "schemaVersion": 1,
        "catalogVersion": CATALOG_VERSION,
        "generatedFrom": {
            "workbook": args.workbook.name,
            "aliasFile": args.aliases.name,
            "classificationSheetCount": category_sheet_count,
            "eligibilityStatus": "pending_post_alias_course_recount",
            "frequencySnapshotUsableForEligibility": False,
        },
        "summary": {
            "mainCategoryCount": len(main_categories),
            "categoryPathCount": len(subcategories),
            "rawTagCount": len(raw_tag_keys),
            "canonicalTagCount": len(serialized_tags),
            "categoryAssignmentCount": assignment_count,
            "rawTagsWithMultipleCategoryPaths": sum(1 for paths in raw_tag_paths.values() if len(paths) > 1),
            "canonicalTagsWithMultipleCategoryPaths": sum(
                1 for item in serialized_tags if len(item["categoryAssignments"]) > 1
            ),
            "approvedAliasCount": len(alias_index),
            "aliasCanonicalTargetsMissingFromCatalog": alias_targets_missing,
        },
        "mainCategories": sorted(main_categories.values(), key=lambda item: (comparison_key(item["name"]), item["id"])),
        "subcategories": sorted(subcategories.values(), key=lambda item: (item["mainCategoryId"], comparison_key(item["name"]), item["id"])),
        "canonicalTags": serialized_tags,
    }

    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(catalog, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(catalog["summary"], ensure_ascii=True, sort_keys=True))


if __name__ == "__main__":
    try:
        main()
    except Exception as error:  # make command-line failures concise and actionable
        print(f"importInterestTagCatalog: {error}", file=sys.stderr)
        raise
